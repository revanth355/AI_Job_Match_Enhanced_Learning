"""Parse user documents in memory. No uploaded files are written to disk."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 60000


def extract_text(data: bytes, filename: str) -> str:
    if not data:
        raise ValueError("The file is empty. Upload a resume with selectable text.")
    if len(data) > MAX_BYTES:
        raise ValueError("The file exceeds the 5 MB limit.")
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".txt":
            text = data.decode("utf-8-sig")
        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported. Use an unlocked copy.")
            if len(reader.pages) > 30:
                raise ValueError("Use a resume with no more than 30 pages.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            from docx import Document
            with ZipFile(BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("This DOCX expands beyond the 20 MB safety limit.")
            doc = Document(BytesIO(data))
            text = "\n".join([p.text for p in doc.paragraphs] + [
                " | ".join(cell.text for cell in row.cells)
                for table in doc.tables for row in table.rows
            ])
        else:
            raise ValueError("Supported formats: PDF, DOCX and UTF-8 TXT.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not read this file. Try exporting it again as PDF, DOCX or UTF-8 TXT.") from exc
    text = text.replace("\x00", "").strip()
    if not text:
        raise ValueError("No readable text found. Scanned PDFs need OCR; paste the text instead.")
    if len(text) > MAX_CHARS:
        raise ValueError(f"The document exceeds {MAX_CHARS:,} characters. Use a shorter resume.")
    return text
