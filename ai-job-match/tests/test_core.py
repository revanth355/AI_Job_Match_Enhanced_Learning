import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import numpy as np
from docx import Document
from pypdf import PdfWriter

from analyzer.core import analyze, find_skills, keyword_similarity, semantic_similarity, markdown_report
from analyzer.extraction import extract_text

ROOT = Path(__file__).resolve().parents[1]


class CoreTests(unittest.TestCase):
    def test_sample_gap_and_denominator(self):
        result = analyze((ROOT / 'samples/resume.txt').read_text(),
                         (ROOT / 'samples/job_description.txt').read_text())
        self.assertEqual(set(result['matched_skills']), {'Python', 'SQL', 'Pandas', 'Git'})
        self.assertEqual(set(result['skills_not_found']),
                         {'Excel', 'Data Visualization', 'Statistics', 'Power BI', 'Communication'})
        self.assertEqual(result['skill_coverage'], 44.4)
        self.assertIn('Python', markdown_report(result))

    def test_word_boundaries_and_aliases(self):
        skills = find_skills('JavaScript and C++ with C#; sklearn, PowerBI. MySQL.')
        self.assertNotIn('Java', skills)
        self.assertNotIn('SQL', skills)
        self.assertEqual(set(skills), {'JavaScript', 'C++', 'C#', 'scikit-learn', 'Power BI', 'MySQL'})

    def test_custom_skills_and_unicode(self):
        skills = find_skills('Ｆｉｇｍａ and PYTHON experience', ['Figma', 'python'])
        self.assertEqual(set(skills), {'Figma', 'Python'})

    def test_no_detected_skills(self):
        result = analyze('Experienced gardener working with flowers and soil for many years.',
                         'Looking for a gardener who can grow flowers and water plants.')
        self.assertIsNone(result['skill_coverage'])

    def test_invalid_inputs(self):
        with self.assertRaises(ValueError):
            analyze('', 'A job requires Python and SQL experience in a team.')
        with self.assertRaises(ValueError):
            analyze('a' * 60001, 'A job requires Python and SQL experience in a team.')

    def test_keyword_extremes(self):
        self.assertAlmostEqual(keyword_similarity('python sql pandas', 'python sql pandas'), 100)
        self.assertEqual(keyword_similarity('python pandas', 'gardening flowers'), 0)
        self.assertEqual(keyword_similarity('the and', 'and the'), 0)

    def test_negation_is_excluded(self):
        self.assertNotIn('Python', find_skills('I have no experience with Python.'))

    def test_semantic_uses_every_token(self):
        class Tokenizer:
            def encode(self, text, add_special_tokens=False):
                return text.split()
            def decode(self, tokens, skip_special_tokens=True):
                return ' '.join(tokens)
        class Model:
            tokenizer = Tokenizer()
            max_seq_length = 6
            def __init__(self):
                self.seen = []
            def encode(self, pieces, **kwargs):
                self.seen.extend(pieces)
                return np.array([[1., 0.] for _ in pieces])
        model = Model()
        self.assertEqual(semantic_similarity('one two three four five six seven', 'eight nine', model), 100)
        self.assertEqual(model.seen, ['one two three four', 'five six seven', 'eight nine'])

    def test_semantic_failure_does_not_silently_fallback(self):
        with patch('analyzer.core.load_model', side_effect=RuntimeError('Unavailable')):
            with self.assertRaises(RuntimeError):
                analyze('Python SQL Pandas data analysis and model experience.',
                        'Python SQL Pandas data analysis and model requirements.', 'semantic')


class ExtractionTests(unittest.TestCase):
    def test_txt(self):
        self.assertEqual(extract_text(b'\xef\xbb\xbfPython resume', 'resume.TXT'), 'Python resume')

    def test_docx_table_and_paragraph(self):
        doc = Document()
        doc.add_paragraph('Python developer')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'SQL'
        stream = BytesIO()
        doc.save(stream)
        text = extract_text(stream.getvalue(), 'resume.docx')
        self.assertIn('Python developer', text)
        self.assertIn('SQL', text)

    def test_pdf_text(self):
        from reportlab.pdfgen.canvas import Canvas
        stream = BytesIO()
        canvas = Canvas(stream)
        canvas.drawString(50, 700, 'Python SQL resume')
        canvas.save()
        self.assertIn('Python SQL resume', extract_text(stream.getvalue(), 'resume.pdf'))

    def test_encrypted_pdf(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt('secret')
        stream = BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, 'Password-protected'):
            extract_text(stream.getvalue(), 'resume.pdf')

    def test_empty_pdf(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        stream = BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, 'Scanned PDFs'):
            extract_text(stream.getvalue(), 'resume.pdf')

    def test_bad_files(self):
        for data, name in [(b'', 'a.txt'), (b'bad', 'a.pdf'), (b'bad', 'a.docx'),
                           (b'bad', 'a.exe'), (b'\xff', 'a.txt'), (b'a' * 5242881, 'a.txt')]:
            with self.subTest(name=name, size=len(data)):
                with self.assertRaises(ValueError):
                    extract_text(data, name)


if __name__ == '__main__':
    unittest.main()
